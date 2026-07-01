"""
Tests for the .docx extractor branch (ADR-044 / unified-portal-strategy.md D5).

Additive fifth extractor branch. These tests assert:
  - _is_docx / detect_type classify a real .docx as DOCX (and do NOT misclassify
    PDF / HTML / legacy .doc / plain zips)
  - extract_docx returns a valid FetchedDocument with text + heading hierarchy
  - route() dispatches a downloaded .docx to extract_docx
  - legacy .doc (OLE) is UNKNOWN → router raises UnsupportedDocTypeError

The .docx payloads are generated in-memory with python-docx, so no binary fixture
is committed.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from src.config.economy_config import EconomyConfig
from src.fetcher.models import FetchedDocument, Zone1Result


def _make_docx_bytes() -> bytes:
    import docx
    d = docx.Document()
    d.add_heading("Privacy Act 1988", level=0)
    d.add_heading("Part I — Preliminary", level=1)
    d.add_paragraph("This is section 1 of the Act about personal information.")
    d.add_heading("Part II — Interpretation", level=1)
    d.add_paragraph("personal information means information about an individual.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Schedule 1"
    table.rows[0].cells[1].text = "Australian Privacy Principles"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _zone1(url: str = "https://www.legislation.gov.au/C2004A03712/latest/word") -> Zone1Result:
    return Zone1Result(
        url=url,
        economy="AU",
        act_title="Privacy Act 1988",
        discovery_tag="KNOWN",
        archive_url="",
    )


def _au_config() -> EconomyConfig:
    return EconomyConfig.model_validate({
        "economy_name": "Australia",
        "iso_code": "AU",
        "un_name": "Australia",
        "script_type": "latin",
        "languages": ["en"],
        "portals": [{"name": "FRL", "url": "https://www.legislation.gov.au"}],
    })


# ── Detection ──────────────────────────────────────────────────────────────────

class TestDocxDetection:
    def test_is_docx_true_for_real_docx(self):
        from src.fetcher.router import _is_docx
        assert _is_docx(_make_docx_bytes()) is True

    def test_is_docx_false_for_pdf(self):
        from src.fetcher.router import _is_docx
        assert _is_docx(b"%PDF-1.7\n...") is False

    def test_is_docx_false_for_plain_zip(self):
        """A ZIP without word/document.xml (e.g. xlsx-like) is not a .docx."""
        from src.fetcher.router import _is_docx
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("xl/workbook.xml", "<x/>")
        assert _is_docx(buf.getvalue()) is False

    def test_detect_type_docx_octet_stream(self):
        """Local .docx downloads arrive as octet-stream — byte sniff must win."""
        from src.fetcher.router import detect_type
        assert detect_type(_make_docx_bytes(), "application/octet-stream") == "DOCX"

    def test_detect_type_docx_proper_content_type(self):
        from src.fetcher.router import detect_type
        ct = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert detect_type(_make_docx_bytes(), ct) == "DOCX"

    def test_detect_type_pdf_not_misclassified_as_docx(self):
        """Regression: the new branch must not disturb PDF detection."""
        from src.fetcher.router import detect_type
        # Minimal PDF header; classify_pdf will call it SCANNED (no text) — the
        # point is it is NOT DOCX.
        assert detect_type(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n", "application/pdf") in ("TEXT_PDF", "SCANNED_PDF")

    def test_detect_type_html_not_misclassified_as_docx(self):
        from src.fetcher.router import detect_type
        html = b"<!doctype html><html><body>" + b"x" * 300 + b"</body></html>"
        assert detect_type(html, "text/html") == "HTML"

    def test_detect_type_legacy_doc_is_unknown(self):
        """Legacy .doc (OLE binary) is NOT a .docx and must fall through to UNKNOWN."""
        from src.fetcher.router import detect_type
        ole_magic = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100
        assert detect_type(ole_magic, "application/msword") == "UNKNOWN"


# ── Extraction ─────────────────────────────────────────────────────────────────

class TestDocxExtraction:
    def test_extract_docx_returns_valid_document(self):
        from src.fetcher.extractors.docx_text import extract_docx
        doc = extract_docx(_make_docx_bytes(), _zone1())
        assert isinstance(doc, FetchedDocument)
        assert doc.doc_type == "DOCX"
        assert doc.extraction_method == "python_docx"
        assert "personal information" in doc.raw_text
        assert doc.cost_log_entry is not None
        assert doc.cost_log_entry.cost_usd == 0.0

    def test_extract_docx_captures_heading_hierarchy(self):
        from src.fetcher.extractors.docx_text import extract_docx
        doc = extract_docx(_make_docx_bytes(), _zone1())
        titles = [s["title"] for s in doc.section_hierarchy]
        assert "Privacy Act 1988" in titles
        assert any("Part I" in t for t in titles)

    def test_extract_docx_includes_table_text(self):
        from src.fetcher.extractors.docx_text import extract_docx
        doc = extract_docx(_make_docx_bytes(), _zone1())
        assert "Australian Privacy Principles" in doc.raw_text

    def test_extract_docx_empty_raises(self):
        from src.fetcher.extractors.docx_text import ExtractionError, extract_docx
        import docx
        buf = io.BytesIO()
        docx.Document().save(buf)  # no content
        with pytest.raises(ExtractionError):
            extract_docx(buf.getvalue(), _zone1())


# ── Router dispatch ──────────────────────────────────────────────────────────────

class TestRouterDocxDispatch:
    def test_route_dispatches_docx_to_extractor(self):
        from src.fetcher import router
        docx_bytes = _make_docx_bytes()

        def mock_download(url, timeout=30):
            return docx_bytes, "application/octet-stream", url

        with pytest.MonkeyPatch.context() as mp:
            mp.setattr(router, "download", mock_download)
            result = router.route(_zone1(), _au_config())

        assert isinstance(result, FetchedDocument)
        assert result.doc_type == "DOCX"
        assert result.extraction_method == "python_docx"
