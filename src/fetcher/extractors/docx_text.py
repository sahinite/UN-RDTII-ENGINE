"""
DOCX (.docx) extraction via python-docx. [ADR-044]

For portals that serve legislation as Word documents (e.g. Australia's FRL offers
a .docx alongside the PDF). Legacy `.doc` (OLE binary) is unreadable by
python-docx, so detect_type classifies it UNKNOWN → the router raises a clear
UnsupportedDocTypeError rather than mis-extracting.
"""

from __future__ import annotations

import io
import time
from typing import TYPE_CHECKING

from src.fetcher.logger import get_logger
from src.fetcher.models import CostLogEntry, FetchedDocument

if TYPE_CHECKING:
    from src.config.economy_config import EconomyConfig
    from src.fetcher.models import Zone1Result

logger = get_logger("docx_extractor")


class ExtractionError(Exception):
    pass


# Word built-in heading styles map to section-hierarchy levels. "Heading 1" → 1,
# "Heading 2" → 2, etc.; "Title" is treated as level 0 (document title).
def _heading_level(style_name: str) -> int | None:
    name = (style_name or "").strip().lower()
    if name == "title":
        return 0
    if name.startswith("heading"):
        tail = name.replace("heading", "").strip()
        if tail.isdigit():
            return int(tail)
        return 1
    return None


def extract_docx(
    raw_bytes: bytes,
    zone1_result: "Zone1Result",
    economy_config: "EconomyConfig | None" = None,
) -> FetchedDocument:
    """Extract text + heading hierarchy from a .docx byte payload."""
    started_at = time.monotonic()

    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guaranteed by requirements
        raise ExtractionError("python-docx not installed") from exc

    try:
        document = docx.Document(io.BytesIO(raw_bytes))
    except Exception as exc:
        raise ExtractionError(
            f"python-docx failed to open document for {zone1_result.url}: {exc}"
        ) from exc

    section_hierarchy: list[dict] = []
    lines: list[str] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        lines.append(text)
        level = _heading_level(para.style.name if para.style else "")
        if level is not None:
            section_hierarchy.append({"level": level, "title": text, "text": ""})

    # Include table cell text (schedules/forms in legislation are often tables).
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    full_text = "\n".join(lines).strip()
    if not full_text:
        raise ExtractionError(f"DOCX contained no extractable text: {zone1_result.url}")

    elapsed_ms = (time.monotonic() - started_at) * 1000

    cost_log = CostLogEntry(
        engine="python_docx",
        pages=None,          # Word documents have no fixed page model
        cost_usd=0.0,        # local/free
        processing_time_ms=elapsed_ms,
    )

    doc = FetchedDocument(
        source_url=zone1_result.url,
        resolved_url=zone1_result.url,
        economy=zone1_result.economy,
        act_title=zone1_result.act_title,
        discovery_tag=zone1_result.discovery_tag,  # type: ignore[arg-type]
        archive_url=zone1_result.archive_url,
        doc_type="DOCX",
        extraction_method="python_docx",
        page_count=None,
        raw_text=full_text,
        section_hierarchy=section_hierarchy,
        cost_log_entry=cost_log,
    )
    doc.validate()

    logger.info({
        "event": "fetch_document_validated",
        "url": zone1_result.url,
        "extraction_method": "python_docx",
        "text_length": len(full_text),
        "sections": len(section_hierarchy),
        "economy": zone1_result.economy,
    })
    return doc
